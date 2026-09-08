"""End-to-end test for scripts/rerender.py: rebuild a finished run's reports from grid.csv.

A grid run costs 12-25 minutes, so this is the escape hatch for an export bug — the test
produces a real run folder with optimize.py on the synthetic cache (reusing test_cli's
_write_cache), deletes the reports, and checks rerender.py puts them back with the same
console block.
"""
from __future__ import annotations

import subprocess
import sys
from pathlib import Path

import pandas as pd
import pytest

from tests.test_cli import SYMBOL, ROOT, _write_cache

RERENDER = [sys.executable, str(ROOT / "scripts" / "rerender.py")]


@pytest.fixture(scope="module")
def finished_run(tmp_path_factory) -> Path:
    """One real optimize.py run folder on the synthetic cache, shared by the tests below."""
    tmp = tmp_path_factory.mktemp("rerender")
    data_dir = tmp / "data"
    data_dir.mkdir()
    _write_cache(data_dir)
    out = tmp / "results" / "rsi2_ema_swing"     # rerender infers the strategy from this parent
    r = subprocess.run([sys.executable, str(ROOT / "scripts" / "optimize.py"),
                        "--strategy", "rsi2_ema_swing", "--tf", "M5",
                        "--axis", "ema=20/100", "--axis", "atr_mult=1.5",
                        "--axis", "rsi_fast=2", "--axis", "rsi2=90/10",
                        "--tp", "4", "--risk", "1",
                        "--data-dir", str(data_dir), "--out", str(out)],
                       cwd=ROOT, capture_output=True, text=True, timeout=600)
    assert r.returncode == 0, r.stdout + r.stderr
    run_dir = next(out.glob("*"))
    (run_dir / "data_dir.txt").write_text(str(data_dir), encoding="utf-8")   # so tests find the cache
    return run_dir


def _data_dir(run_dir: Path) -> str:
    return (run_dir / "data_dir.txt").read_text(encoding="utf-8")


def test_rerender_rebuilds_the_reports_in_place(finished_run):
    xlsx, html = finished_run / f"report_{SYMBOL}.xlsx", finished_run / f"report_{SYMBOL}.html"
    assert xlsx.exists() and html.exists()
    grid_before = (finished_run / "grid.csv").read_bytes()
    xlsx.unlink()
    html.unlink()

    r = subprocess.run(RERENDER + [str(finished_run), "--symbol", SYMBOL,
                                   "--data-dir", _data_dir(finished_run)],
                       cwd=ROOT, capture_output=True, text=True, timeout=600)
    assert r.returncode == 0, r.stdout + r.stderr
    assert xlsx.exists() and html.exists()
    assert (finished_run / "grid.csv").read_bytes() == grid_before      # grid.csv is read, never rewritten

    # The console block a crashed run would have lost.
    assert "=== Recommendation ===" in r.stdout
    assert "top 5 by IS avg R" in r.stdout
    text = html.read_text(encoding="utf-8")
    assert "rsi2_ema_swing" in text and "Recommendation" in text        # strategy inferred from the folder

    # Re-running over its own output is idempotent (overwrite, not a second folder).
    r2 = subprocess.run(RERENDER + [str(finished_run), "--symbol", SYMBOL,
                                    "--data-dir", _data_dir(finished_run)],
                        cwd=ROOT, capture_output=True, text=True, timeout=600)
    assert r2.returncode == 0, r2.stdout + r2.stderr
    assert {p.name for p in finished_run.iterdir()} >= {"grid.csv", f"report_{SYMBOL}.xlsx",
                                                        f"report_{SYMBOL}.html"}


def test_rerender_defaults_the_run_switches_from_the_first_grid_row(finished_run):
    """grid.csv records min_sl_mult / htf_seconds per row; the run used 0/0 and rerender says so."""
    grid = pd.read_csv(finished_run / "grid.csv")
    assert (grid["min_sl_mult"] == 0.0).all() and (grid["htf_seconds"] == 0).all()
    r = subprocess.run(RERENDER + [str(finished_run), "--symbol", SYMBOL,
                                   "--data-dir", _data_dir(finished_run)],
                       cwd=ROOT, capture_output=True, text=True, timeout=600)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "min_sl_mult=0, htf_seconds=0" in r.stdout
    assert "--help" not in r.stdout


def test_rerender_reruns_the_recommended_combo_and_writes_its_sheets(finished_run, tmp_path):
    """The path that actually costs money to get wrong: a grid row that clears recommend()'s
    gates must be re-run with one run_single, producing Trades_/Monthly_/Equity_ sheets.

    Random-walk synthetic bars never clear the profit gates, so the row is doctored to pass
    (the numbers are irrelevant — what is tested is that rerender reconstructs axis_values from
    rec[tf]["params"], loads the bars and re-runs the combo).
    """
    import openpyxl

    run_dir = tmp_path / "results" / "rsi2_ema_swing" / "20260908_000000"
    run_dir.mkdir(parents=True)
    grid = pd.read_csv(finished_run / "grid.csv")
    grid.loc[:, ["is_n_trades", "n_trades"]] = 40
    grid.loc[:, "oos_n_trades"] = 15
    grid.loc[:, ["is_avg_r", "oos_avg_r", "avg_r", "robust_r"]] = 0.3
    grid.loc[:, "net_pnl"] = 1_500.0
    grid.loc[:, "profit_factor"] = 1.4
    grid.loc[:, "max_dd_pct"] = 0.1
    grid.loc[:, ["oversized_share", "capped_share"]] = 0.0
    grid.loc[:, "ruined"] = False
    grid.to_csv(run_dir / "grid.csv", index=False)

    r = subprocess.run(RERENDER + [str(run_dir), "--symbol", SYMBOL,
                                   "--data-dir", _data_dir(finished_run)],
                       cwd=ROOT, capture_output=True, text=True, timeout=600)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "no reliable parameter set" not in r.stdout
    # params rebuilt from the picked row, printed exactly as optimize.py prints them
    assert "M5: rsi_fast 2 f_hi 90 f_lo 10 ema_fast 20 ema_slow 100 atr_mult 1.5 tp_r 4" in r.stdout
    assert "bars usable" in r.stdout                           # the bars were actually loaded

    wb = openpyxl.load_workbook(run_dir / f"report_{SYMBOL}.xlsx", read_only=True)
    assert {"Summary", "Grid", "Params", "Trades_M5", "Monthly_M5", "Equity_M5"} <= set(wb.sheetnames)
    assert wb["Trades_M5"].max_row > 1                         # run_single produced real trades
    html = (run_dir / f"report_{SYMBOL}.html").read_text(encoding="utf-8")
    assert "EMA 20/100" in html and "Equity" in html


def test_rerender_help_documents_the_first_row_defaults():
    r = subprocess.run(RERENDER + ["--help"], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert r.returncode == 0
    assert "first row" in r.stdout                     # --min-sl-mult / --htf say where they default from


def test_rerender_rejects_a_folder_without_a_grid(tmp_path):
    empty = tmp_path / "results" / "rsi2_ema_swing" / "20260101_000000"
    empty.mkdir(parents=True)
    r = subprocess.run(RERENDER + [str(empty)], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert r.returncode != 0
    combined = r.stdout + r.stderr
    assert "no grid.csv" in combined and "Traceback" not in combined


def test_rerender_rejects_an_unrecognisable_strategy_folder(tmp_path):
    bogus = tmp_path / "somewhere" / "20260101_000000"
    bogus.mkdir(parents=True)
    (bogus / "grid.csv").write_text("tf,tp_r\nM5,2.0\n", encoding="utf-8")
    r = subprocess.run(RERENDER + [str(bogus)], cwd=ROOT, capture_output=True, text=True, timeout=120)
    assert r.returncode != 0
    combined = r.stdout + r.stderr
    assert "somewhere" in combined and "Traceback" not in combined
