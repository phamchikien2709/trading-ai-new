from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

from rsi_fvg.backtest.export import write_csvs, write_html, write_xlsx
from rsi_fvg.backtest.optimize import GridSpec, recommend, run_optimization, run_single
from rsi_fvg.bars import Bars
from rsi_fvg.params import CostParams, SizingParams, SymbolSpec
from rsi_fvg.strategies.rsi2_swing import Rsi2SwingParams

SPEC = SymbolSpec(name="T", point=0.01, digits=2, contract_size=1.0)
COSTS = CostParams(spread_points=20, commission_per_lot_rt=0.0, slippage_points=0)
SIZING = SizingParams(risk_pct=5.0, initial_equity=10_000.0)
SMALL = GridSpec(tp_r=(1.0, 2.0), atr_mult=(0.5, 1.0), rsi_slow_levels=((75.0, 25.0),), rsi_fast_levels=((90.0, 10.0),))


def _bars(n=6000, seed=5):
    rng = np.random.default_rng(seed)
    close = 2000 + np.cumsum(rng.normal(0, 2, n))
    open_ = np.r_[close[0], close[:-1]]
    high = np.maximum(open_, close) + rng.uniform(0.1, 1.5, n)
    low = np.minimum(open_, close) - rng.uniform(0.1, 1.5, n)
    return Bars(time=1_700_000_000 + np.arange(n, dtype=np.int64) * 300, open=open_, high=high, low=low, close=close)


def _fixture():
    b = _bars()
    df = run_optimization({"M5": b}, {"M5": SPEC}, Rsi2SwingParams(), SMALL, COSTS, SIZING)
    rec = recommend(df)
    row = df.iloc[0]
    _, res = run_single(b, SPEC, Rsi2SwingParams(atr_mult=row.atr_mult), row.tp_r, COSTS, SIZING, "hedge")
    info = {"symbol": "T", "timeframes": ["M5"], "data_range": {"M5": ("2023-11-14", "2023-12-05")},
            "initial_equity": 10_000, "risk_pct": 5.0, "concurrency": "hedge", "spread_points": 20,
            "commission_per_lot_rt": 0.0, "slippage_points": 0, "is_frac": 0.7,
            "grid": {"tp_r": [1, 2], "atr_mult": [0.5, 1], "rsi14": ["75/25"], "rsi2": ["90/10"]},
            "git_hash": "test", "generated_at": "2026-09-08 00:00"}
    return df, rec, {"M5": res}, info


def test_write_csvs(tmp_path):
    df, rec, res, info = _fixture()
    write_csvs(tmp_path, df, res)
    assert (tmp_path / "grid.csv").exists() and (tmp_path / "trades_M5.csv").exists()
    assert len(pd.read_csv(tmp_path / "grid.csv")) == 4


def test_write_xlsx_sheets_and_rows(tmp_path):
    df, rec, res, info = _fixture()
    p = tmp_path / "r.xlsx"
    write_xlsx(p, df, rec, res, info)
    wb = openpyxl.load_workbook(p, read_only=True)
    assert {"Summary", "Grid", "Trades_M5", "Monthly_M5", "Equity_M5", "Params"} <= set(wb.sheetnames)
    assert wb["Grid"].max_row == 5            # header + 4 combos
    head = [c.value for c in next(wb["Grid"].iter_rows(min_row=1, max_row=1))]
    assert head[:7] == ["tf", "ob", "os", "f_hi", "f_lo", "atr_mult", "tp_r"]
    assert wb["Equity_M5"].max_row <= 20_001   # downsampled


def test_write_xlsx_handles_no_recommendation(tmp_path):
    df, _, res, info = _fixture()
    write_xlsx(tmp_path / "r.xlsx", df, {"M5": None}, res, info)
    wb = openpyxl.load_workbook(tmp_path / "r.xlsx", read_only=True)
    cells = [c.value for row in wb["Summary"].iter_rows() for c in row if isinstance(c.value, str)]
    assert any("no reliable" in v for v in cells)


def test_write_html_contents(tmp_path):
    df, rec, res, info = _fixture()
    p = tmp_path / "r.html"
    write_html(p, df, rec, res, info)
    html = p.read_text(encoding="utf-8")
    assert "Recommendation" in html and "plotly" in html.lower()
    assert "Heatmap" in html or "heatmap" in html
    assert "cdn.plot.ly" in html                   # cdn mode by default
    write_html(tmp_path / "off.html", df, rec, res, info, offline=True)
    assert (tmp_path / "off.html").stat().st_size > 1_000_000   # plotly.js embedded


def test_write_html_handles_incomplete_grid(tmp_path):
    df, rec, res, info = _fixture()
    incomplete_df = df.iloc[:-1]  # Drop one row to create incomplete (3-of-4) grid
    p = tmp_path / "r.html"
    write_html(p, incomplete_df, rec, res, info)
    html = p.read_text(encoding="utf-8")
    assert p.exists()
    assert "Recommendation" in html
    # Also test write_xlsx does not crash on incomplete grid
    xlsx_p = tmp_path / "r.xlsx"
    write_xlsx(xlsx_p, incomplete_df, rec, res, info)
    assert xlsx_p.exists()
