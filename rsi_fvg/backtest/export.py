"""Excel workbook + self-contained HTML report for the optimisation run (spec §5)."""
from __future__ import annotations

import html as _html
from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from plotly.subplots import make_subplots

from ..strategies.registry import StrategyAdapter, get_adapter
from .engine import BacktestResult
from .metrics import monthly_table
from .optimize import FILTER_TEXT

GRID_METRIC_COLS = ["n_signals", "n_rejected_min_sl",
                    "n_trades", "win_rate", "avg_r", "profit_factor", "max_dd_pct",
                    "net_pnl", "final_equity", "ruined", "oversized_share", "capped_share",
                    "is_n_trades", "is_avg_r", "is_max_dd_pct", "oos_n_trades", "oos_avg_r",
                    "oos_max_dd_pct", "robust_r", "robust_ratio", "grid_edge", "flags"]


def grid_first_cols(adapter: StrategyAdapter) -> list[str]:
    return adapter.full_key_cols() + GRID_METRIC_COLS


def adapter_from_run_info(run_info: dict) -> StrategyAdapter:
    return get_adapter(run_info.get("strategy", "rsi2_swing"))


GRID_FIRST_COLS = grid_first_cols(get_adapter("rsi2_swing"))
REC_COLS = ["n_trades", "win_rate", "avg_r", "profit_factor", "max_dd_pct", "net_pnl", "ruined",
            "oversized_share", "capped_share", "is_n_trades", "is_avg_r", "oos_n_trades", "oos_avg_r",
            "robust_r", "robust_ratio", "grid_edge"]
FLAG_NAMES = "ruined, n&lt;30, oos_sign_flip, buy_sell_imbalance, oversized, capped, grid_edge"
EQUITY_MAX_ROWS = 20_000
_RED, _WHITE, _GREEN = "F8696B", "FFFFFF", "63BE7B"

# ---------------------------------------------------------------- palette ----
# Validated data-viz palette (dataviz skill, references/palette.md), applied to
# the plotly figures below: a blue<->red diverging pair with a neutral gray
# midpoint (never a hue at the midpoint), a single-hue blue sequential ramp,
# and chart chrome (surface/ink/gridline) tokens for a consistent, muted look.
_INK = "#0b0b0b"
_MUTED = "#898781"
_GRID = "#e1e0d9"
_AXIS_LINE = "#c3c2b7"
_SURFACE = "#fcfcfb"
_SERIES_BLUE = "#2a78d6"
_SERIES_RED = "#e34948"
_NEUTRAL_MID = "#f0efec"
_FONT_FAMILY = "system-ui, -apple-system, 'Segoe UI', sans-serif"
_DIVERGING = [[0.0, _SERIES_RED], [0.5, _NEUTRAL_MID], [1.0, _SERIES_BLUE]]
_SEQ_BLUE = [[0.0, "#cde2fb"], [0.2, "#9ec5f4"], [0.4, "#6da7ec"], [0.6, "#3987e5"], [0.8, "#256abf"], [1.0, "#0d366b"]]


def _apply_theme(fig: go.Figure) -> go.Figure:
    fig.update_layout(paper_bgcolor=_SURFACE, plot_bgcolor=_SURFACE,
                      font=dict(family=_FONT_FAMILY, color=_INK, size=12),
                      title_font=dict(size=15, color=_INK))
    fig.update_xaxes(gridcolor=_GRID, linecolor=_AXIS_LINE, zerolinecolor=_AXIS_LINE, tickfont=dict(color=_MUTED))
    fig.update_yaxes(gridcolor=_GRID, linecolor=_AXIS_LINE, zerolinecolor=_AXIS_LINE, tickfont=dict(color=_MUTED))
    return fig


def _order_grid(df: pd.DataFrame, adapter: StrategyAdapter) -> pd.DataFrame:
    first = [c for c in grid_first_cols(adapter) if c in df.columns]
    rest = [c for c in df.columns if c not in first]
    return df[first + rest]


def _strip_tz(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    for c in out.columns:
        if isinstance(out[c].dtype, pd.DatetimeTZDtype):
            out[c] = out[c].dt.tz_localize(None)
    return out


def _downsample(obj, max_rows: int = EQUITY_MAX_ROWS):
    step = max(1, int(np.ceil(len(obj) / max_rows)))
    return obj.iloc[::step]


def _drawdown_frame(eq: pd.Series) -> pd.DataFrame:
    """equity + running drawdown, computed on the FULL curve.

    Downsampling first drops the very bars that make the peaks and troughs, so the reported
    drawdown came out shallower than the run's real one — take the drawdown first, thin the
    result afterwards.
    """
    peak = eq.cummax()
    dd = peak - eq
    return pd.DataFrame({"equity": eq, "drawdown_usd": dd,
                         "drawdown_pct": (dd / peak.replace(0, np.nan)).fillna(0.0) * 100.0})


def _equity_frame(res: BacktestResult, max_rows: int = EQUITY_MAX_ROWS) -> pd.DataFrame:
    d = _downsample(_drawdown_frame(res.equity), max_rows)
    idx = d.index.tz_localize(None) if d.index.tz is not None else d.index
    return pd.DataFrame({"time": idx, "equity": d["equity"].values,
                         "drawdown_usd": d["drawdown_usd"].values, "drawdown_pct": d["drawdown_pct"].values})


def _rec_table(rec: dict) -> pd.DataFrame:
    rows = []
    for tf, r in rec.items():
        if r is None:
            rows.append({"tf": tf, "status": f"no reliable parameter set (filters: {FILTER_TEXT})"})
        else:
            row = {"tf": tf, "status": "recommended"}
            row.update({k: v for k, v in r["params"].items() if k != "tf"})
            row.update({k: r["row"].get(k) for k in REC_COLS})
            row["reason"] = r["reason"]
            rows.append(row)
    return pd.DataFrame(rows)


def _info_table(run_info: dict) -> pd.DataFrame:
    flat = []
    for k, v in run_info.items():
        if k == "grid":
            for gk, gv in v.items():
                flat.append((f"grid.{gk}", ", ".join(str(x) for x in gv)))
        elif k == "data_range":
            for tf, (a, b) in v.items():
                flat.append((f"data.{tf}", f"{a} -> {b}"))
        elif isinstance(v, (list, tuple)):
            flat.append((k, ", ".join(str(x) for x in v)))
        else:
            flat.append((k, v))
    return pd.DataFrame(flat, columns=["key", "value"])


def write_csvs(out_dir: Path, grid_df: pd.DataFrame, rec_results: dict[str, BacktestResult],
               run_info: dict) -> None:
    adapter = adapter_from_run_info(run_info)
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    _order_grid(grid_df, adapter).to_csv(out_dir / "grid.csv", index=False)
    for tf, res in rec_results.items():
        res.trades.to_csv(out_dir / f"trades_{tf}.csv", index=False)
        if len(res.skipped):
            res.skipped.to_csv(out_dir / f"skipped_{tf}.csv", index=False)


def _color_scale(ws, col_letter: str, first_row: int, last_row: int) -> None:
    if last_row < first_row:
        return
    ws.conditional_formatting.add(
        f"{col_letter}{first_row}:{col_letter}{last_row}",
        ColorScaleRule(start_type="min", start_color=_RED, mid_type="num", mid_value=0, mid_color=_WHITE,
                       end_type="max", end_color=_GREEN))


def _bold_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.font = Font(bold=True)


def write_xlsx(path: Path, grid_df: pd.DataFrame, rec: dict, rec_results: dict[str, BacktestResult],
               run_info: dict) -> None:
    adapter = adapter_from_run_info(run_info)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    grid = _order_grid(grid_df.copy(), adapter)
    for col in ("split_time", "ruin_time"):
        if col in grid.columns:
            grid[col] = pd.to_datetime(grid[col], utc=True).dt.tz_localize(None)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        info = _info_table(run_info)
        info.to_excel(xw, sheet_name="Summary", index=False, startrow=0)
        rec_tbl = _rec_table(rec)
        rec_tbl.to_excel(xw, sheet_name="Summary", index=False, startrow=len(info) + 3)
        ws = xw.sheets["Summary"]
        ws.cell(row=len(info) + 3, column=1, value="Recommendation per timeframe").font = Font(bold=True)
        _bold_header(ws, 1)
        _bold_header(ws, len(info) + 4)
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 60

        grid.to_excel(xw, sheet_name="Grid", index=False)
        wg = xw.sheets["Grid"]
        wg.freeze_panes = "A2"
        wg.auto_filter.ref = wg.dimensions
        _bold_header(wg)
        for col in ("oos_avg_r", "is_avg_r", "avg_r", "robust_r"):
            if col in grid.columns:
                _color_scale(wg, get_column_letter(grid.columns.get_loc(col) + 1), 2, len(grid) + 1)

        for tf, res in rec_results.items():
            _strip_tz(res.trades).to_excel(xw, sheet_name=f"Trades_{tf}", index=False)
            _bold_header(xw.sheets[f"Trades_{tf}"])
            xw.sheets[f"Trades_{tf}"].freeze_panes = "A2"
            mt = monthly_table(res.equity)
            mt.to_excel(xw, sheet_name=f"Monthly_{tf}")
            if not mt.empty:
                _color_scale_block(xw.sheets[f"Monthly_{tf}"], 2, len(mt) + 1, 2, 13)
            _equity_frame(res).to_excel(xw, sheet_name=f"Equity_{tf}", index=False)
            _bold_header(xw.sheets[f"Equity_{tf}"])

        params = _info_table({k: run_info[k] for k in ("grid", "initial_equity", "risk_pct", "concurrency",
                                                        "spread_points", "commission_per_lot_rt",
                                                        "slippage_points", "is_frac", "min_sl_mult",
                                                        "htf_seconds", "git_hash") if k in run_info})
        params.to_excel(xw, sheet_name="Params", index=False)
        _bold_header(xw.sheets["Params"])


def _color_scale_block(ws, r1: int, r2: int, c1: int, c2: int) -> None:
    ws.conditional_formatting.add(
        f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}",
        ColorScaleRule(start_type="min", start_color=_RED, mid_type="num", mid_value=0, mid_color=_WHITE,
                       end_type="max", end_color=_GREEN))


# ------------------------------------------------------------------ HTML ----
_CSS = """
body{font-family:system-ui,-apple-system,"Segoe UI",Arial,sans-serif;margin:24px;color:#0b0b0b;background:#f9f9f7}
h1{font-size:22px} h2{font-size:18px;margin-top:36px;border-bottom:1px solid #e1e0d9;padding-bottom:4px}
table{border-collapse:collapse;font-size:13px} th,td{border:1px solid #e1e0d9;padding:4px 8px;text-align:right}
th{background:#f0efec} td:first-child,th:first-child{text-align:left}
.rec{background:#fcfcfb;border:1px solid #e1e0d9;border-radius:6px;padding:12px 16px;margin:8px 0}
.rec.none{border-color:#fab219} .warn{color:#d03b3b} .muted{color:#52514e;font-size:12px}
"""


def _fig_equity(tf: str, res: BacktestResult) -> go.Figure:
    d = _downsample(_drawdown_frame(res.equity), 5000)      # drawdown from the full curve (M3)
    fig = make_subplots(rows=2, cols=1, shared_xaxes=True, row_heights=[0.7, 0.3], vertical_spacing=0.04)
    fig.add_trace(go.Scatter(x=d.index, y=d["equity"].values, name="Equity",
                             line=dict(width=2, color=_SERIES_BLUE)), row=1, col=1)
    fig.add_trace(go.Scatter(x=d.index, y=-d["drawdown_pct"].values, name="Drawdown %", fill="tozeroy",
                             line=dict(width=2, color=_SERIES_RED),
                             fillcolor="rgba(227,73,72,0.12)"), row=2, col=1)
    ruin = " · RUINED" if res.ruined else ""
    fig.update_layout(title=f"{tf} — Equity & drawdown (recommended parameters){ruin}", height=520,
                      margin=dict(l=40, r=20, t=50, b=30), legend=dict(orientation="h"))
    fig.update_yaxes(title_text="Equity (USD)", row=1, col=1)
    fig.update_yaxes(title_text="Drawdown (%)", row=2, col=1)
    return _apply_theme(fig)


def _fig_r_hist(tf: str, res: BacktestResult) -> go.Figure:
    fig = go.Figure(go.Histogram(x=res.trades["r_multiple"], nbinsx=40, marker_color=_SERIES_BLUE))
    fig.update_layout(title=f"{tf} — Distribution of R per trade", height=320, bargap=0.05,
                      margin=dict(l=40, r=20, t=50, b=30), xaxis_title="R multiple", yaxis_title="Trade count")
    return _apply_theme(fig)


def _fig_heatmaps(tf: str, g: pd.DataFrame, adapter: StrategyAdapter) -> go.Figure:
    y_col = adapter.axis(adapter.robust_axis).columns[0]
    panel_cols = adapter.panel_cols
    groups = list(g.groupby(panel_cols, sort=True))
    n = len(groups)
    cols = min(3, max(n, 1))
    rows = int(np.ceil(n / cols)) if n else 1
    titles = [adapter.title(dict(zip(panel_cols, key if isinstance(key, tuple) else (key,))))
              for key, _ in groups]
    fig = make_subplots(rows=rows, cols=cols, subplot_titles=titles,
                        horizontal_spacing=0.06, vertical_spacing=0.12)
    zmax = float(np.nanmax(np.abs(g["oos_avg_r"].to_numpy()))) if len(g) else 1.0
    zmax = max(zmax, 1e-9)
    for i, (_, gg) in enumerate(groups):
        piv = gg.pivot(index=y_col, columns="tp_r", values="oos_avg_r").sort_index()
        ntr = gg.pivot(index=y_col, columns="tp_r", values="n_trades").reindex_like(piv)
        text = [[("" if (pd.isna(v) or pd.isna(k)) else f"{v:+.2f}<br>n={int(k)}")
                 for v, k in zip(rv, rk)] for rv, rk in zip(piv.values, ntr.values)]
        fig.add_trace(go.Heatmap(z=piv.values, x=[f"TP {c:g}R" for c in piv.columns],
                                 y=[f"{y_col}×{r:g}" for r in piv.index], colorscale=_DIVERGING, zmid=0,
                                 zmin=-zmax, zmax=zmax, text=text, texttemplate="%{text}",
                                 showscale=(i == 0), colorbar=dict(title="OOS avg R")),
                      row=i // cols + 1, col=i % cols + 1)
    fig.update_layout(title=f"{tf} — OOS avg R heatmap (TP × {y_col}) per parameter set",
                      height=300 * rows + 80, margin=dict(l=40, r=20, t=70, b=30))
    return _apply_theme(fig)


def _fig_is_oos(tf: str, g: pd.DataFrame, adapter: StrategyAdapter) -> go.Figure:
    labels = [f"TP {row['tp_r']:g}R · {adapter.title(row)}"
              f"<br>n={int(row['n_trades'])} · flags: {row['flags'] or '-'}"
              for row in g.to_dict("records")]
    fig = go.Figure(go.Scatter(x=g["is_avg_r"], y=g["oos_avg_r"], mode="markers", text=labels,
                               hovertemplate="%{text}<br>IS %{x:+.2f}R · OOS %{y:+.2f}R<extra></extra>",
                               marker=dict(size=8, color=g["n_trades"], colorscale=_SEQ_BLUE, showscale=True,
                                           colorbar=dict(title="n trades"),
                                           line=dict(width=0.5, color=_AXIS_LINE))))
    lim = float(np.nanmax(np.abs(np.r_[g["is_avg_r"].to_numpy(), g["oos_avg_r"].to_numpy()]))) if len(g) else 1.0
    lim = max(lim, 0.1) * 1.1
    fig.add_shape(type="line", x0=-lim, y0=-lim, x1=lim, y1=lim, line=dict(dash="dot", color=_MUTED))
    fig.add_hline(y=0, line=dict(color=_GRID, width=1))
    fig.add_vline(x=0, line=dict(color=_GRID, width=1))
    fig.update_layout(title=f"{tf} — IS vs OOS avg R (each point = one combo)", height=420,
                      xaxis_title="IS avg R (R multiples)", yaxis_title="OOS avg R (R multiples)",
                      margin=dict(l=40, r=20, t=50, b=30))
    return _apply_theme(fig)


def _fmt_table(df: pd.DataFrame, cols: list[str]) -> str:
    d = df[cols].copy()
    for c in cols:
        if pd.api.types.is_float_dtype(d[c]):
            d[c] = d[c].map(lambda v: f"{v:.3f}" if pd.notna(v) else "")
    return d.to_html(index=False, escape=True, border=0)


def write_html(path: Path, grid_df: pd.DataFrame, rec: dict, rec_results: dict[str, BacktestResult],
               run_info: dict, offline: bool = False) -> None:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    adapter = adapter_from_run_info(run_info)
    include = True if offline else "cdn"
    state = {"first": True}

    def fig_html(fig: go.Figure) -> str:
        js = include if state["first"] else False
        state["first"] = False
        return fig.to_html(full_html=False, include_plotlyjs=js)

    e = _html.escape
    parts = [f"<h1>RSI2 Swing Pullback — Backtest & Optimisation — {e(str(run_info.get('symbol', '')))}</h1>",
             f"<p class='muted'>Generated {e(str(run_info.get('generated_at', '')))} · code {e(str(run_info.get('git_hash', '')))}</p>",
             _fmt_table(_info_table(run_info), ["key", "value"])]

    parts.append("<h2>Recommendation</h2>")
    for tf, r in rec.items():
        if r is None:
            parts.append(f"<div class='rec none'><b>{e(tf)}</b>: no reliable parameter set. "
                         f"No combo passed: {e(FILTER_TEXT)}.</div>")
        else:
            p, row = r["params"], r["row"]
            parts.append(f"<div class='rec'><b>{e(tf)}</b>: TP <b>{p['tp_r']:g}R</b>, ATR mult <b>{p['atr_mult']:g}</b>, "
                         f"RSI14 <b>{p['ob']:g}/{p['os']:g}</b>, RSI fast <b>{p['rsi_fast']:g}</b> "
                         f"<b>{p['f_hi']:g}/{p['f_lo']:g}</b>"
                         f"<br>net P&amp;L <b>${row['net_pnl']:,.0f}</b> · profit factor <b>{row['profit_factor']:.2f}</b> · "
                         f"max drawdown <b>{row['max_dd_pct']:.1%}</b> · capped lots "
                         f"<b>{row.get('capped_share', 0.0):.0%}</b>"
                         f"<br><span class='muted'>{e(r['reason'])}</span></div>")

    for tf, res in rec_results.items():
        parts.append(f"<h2>{e(tf)} — recommended parameters</h2>")
        parts.append(fig_html(_fig_equity(tf, res)))
        parts.append(fig_html(_fig_r_hist(tf, res)))

    for tf, g in grid_df.groupby("tf", sort=False):
        parts.append(f"<h2>{e(tf)} — parameter grid</h2>")
        parts.append(fig_html(_fig_heatmaps(tf, g, adapter)))
        parts.append(fig_html(_fig_is_oos(tf, g, adapter)))
        # Ranked on what recommend() actually scores (IS + robustness); OOS is only a gate.
        top = g.sort_values(["is_avg_r", "robust_r"], ascending=False).head(10)
        parts.append("<h3>Top 10 by IS avg R (then robustness) — the ranking the recommendation uses</h3>")
        parts.append(_fmt_table(top, [c for c in grid_first_cols(adapter) if c in top.columns]))

    parts.append("<h2>Warnings</h2>")
    flagged = grid_df[grid_df["flags"].astype(str) != ""]
    if flagged.empty:
        parts.append("<p>none</p>")
    else:
        parts.append(f"<p class='warn'>{len(flagged)} of {len(grid_df)} combos carry flags "
                     f"({FLAG_NAMES}). See the Grid sheet / grid.csv.</p>")
        counts = flagged["flags"].str.split(";").explode().value_counts()
        parts.append(_fmt_table(counts.rename_axis("flag").reset_index(name="combos"), ["flag", "combos"]))

    html = ("<!doctype html><html><head><meta charset='utf-8'><title>RSI2 Swing report</title>"
            f"<style>{_CSS}</style></head><body>{''.join(parts)}</body></html>")
    path.write_text(html, encoding="utf-8")
